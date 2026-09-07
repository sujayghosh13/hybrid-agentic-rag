import io
from pathlib import Path
import docx
import pytest

from src.ingestion.loaders.docx_loader import DocxLoader
from src.ingestion.loaders.factory import get_loader
from src.ingestion.models import DocType
from src.ingestion.pipeline import IngestionPipeline


def build_test_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    """Helper to create an in-memory DOCX file.
    
    Args:
        paragraphs: List of (style_name, text) tuples.
    """
    doc = docx.Document()
    for style, text in paragraphs:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_loader_can_load():
    loader = DocxLoader()
    assert loader.can_load(Path("test.docx")) is True
    assert loader.can_load(Path("test.doc")) is True
    assert loader.can_load(Path("test.pdf")) is False
    assert loader.can_load(Path("test.html")) is False


def test_factory_returns_docx_loader():
    loader = get_loader(Path("sample_architecture.docx"))
    assert isinstance(loader, DocxLoader)


def test_docx_loader_parse_headings_and_tables(tmp_path):
    doc_path = tmp_path / "k8s_storage.docx"
    doc = docx.Document()
    doc.add_heading("Kubernetes Persistent Volumes", level=1)
    doc.add_paragraph("A PersistentVolume (PV) is a piece of storage in the cluster that has been provisioned by an administrator.")
    doc.add_heading("Access Modes", level=2)
    doc.add_paragraph("A PersistentVolume can be mounted on a host in any way supported by the resource provider.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Mode"
    table.rows[0].cells[1].text = "Description"
    table.rows[1].cells[0].text = "ReadWriteOnce"
    table.rows[1].cells[1].text = "Mounted as read-write by a single node"
    
    doc.save(str(doc_path))

    loader = DocxLoader()
    raw_doc = loader.load(doc_path)

    assert raw_doc.doc_type == DocType.DOCX
    assert raw_doc.filename == "k8s_storage.docx"
    assert len(raw_doc.elements) >= 3

    # Check headings and hierarchy
    headings = [el.heading for el in raw_doc.elements if el.heading]
    assert "Kubernetes Persistent Volumes" in headings or "Access Modes" in headings
    
    # Check table content extracted
    all_text = " ".join(el.text for el in raw_doc.elements)
    assert "ReadWriteOnce" in all_text
    assert "PersistentVolume (PV)" in all_text


def test_docx_ingestion_pipeline_chunking(tmp_path):
    doc_path = tmp_path / "microservices.docx"
    content = build_test_docx([
        ("Heading 1", "Microservices Networking Specification"),
        ("Normal", "Each microservice runs in an isolated container network namespace."),
        ("Heading 2", "Service Discovery"),
        ("Normal", "Internal DNS resolves service names to virtual IP addresses within the bridge network."),
    ])
    doc_path.write_bytes(content)

    pipeline = IngestionPipeline()
    chunks = pipeline.process_file(doc_path)

    assert len(chunks) >= 1
    assert chunks[0].metadata.doc_type == "docx"
    assert chunks[0].metadata.filename == "microservices.docx"
    assert any("Service Discovery" in (c.metadata.section or "") or "DNS" in c.text for c in chunks)


def test_docx_empty_file_handling(tmp_path):
    doc_path = tmp_path / "empty.docx"
    doc = docx.Document()
    doc.save(str(doc_path))

    loader = DocxLoader()
    raw_doc = loader.load(doc_path)
    assert len(raw_doc.elements) == 0


def test_docx_file_not_found():
    loader = DocxLoader()
    with pytest.raises(FileNotFoundError):
        loader.load(Path("non_existent_file.docx"))
